"""Unit tests for in-place `.docx` formatting preservation, C2PA stripping, and stateless ZIP bundling."""

from __future__ import annotations

import io
import zipfile

import docx

from draftproof.config import ConfigNamespace
from draftproof.core.doc_processor import (
    create_checker_reports_zip,
    create_citations_zip,
    create_humanized_docs_zip,
    process_docx_bytes,
)
from draftproof.core.modifier import TextModifier


def test_docx_in_place_run_formatting_and_c2pa_stripping() -> None:
    """Bold/italic runs and tables must be preserved while prose is humanized and metadata stripped."""
    doc = docx.Document()
    doc.core_properties.author = "Tracked AI Author"
    doc.core_properties.comments = "C2PA_MANIFEST_DATA"

    p = doc.add_paragraph()
    r1 = p.add_run("In conclusion, it is important to delve into ")
    r1.bold = True
    r2 = p.add_run("the comprehensive tapestry of neural networks.")
    r2.italic = True

    tbl = doc.add_table(rows=1, cols=1)
    tbl.cell(0, 0).paragraphs[0].text = (
        "Furthermore, this approach plays a pivotal role in providing a robust framework for analysis."
    )

    buf = io.BytesIO()
    doc.save(buf)
    raw_bytes = buf.getvalue()

    cfg = ConfigNamespace({"style_preset": "conversational", "intensity": 80})
    cfg.apply_intensity(80)
    mod = TextModifier(config=cfg)

    res = process_docx_bytes(
        raw_bytes,
        lambda txt, sec=None: mod.modify_line(txt, section_name=sec or "body"),
        filename="sample.docx",
        strip_c2pa=True,
    )
    assert res.paragraphs_modified >= 1
    assert res.c2pa_stripped is True

    out_doc = docx.Document(io.BytesIO(res.output_bytes))
    assert out_doc.core_properties.author == ""
    assert out_doc.core_properties.comments == ""
    # Verify runs and bold/italic attributes survived
    out_runs = out_doc.paragraphs[0].runs
    assert len(out_runs) == 2
    assert out_runs[0].bold is True
    assert out_runs[1].italic is True
    full_after = out_doc.paragraphs[0].text
    assert "delve into" not in full_after.lower()
    assert "tapestry" not in full_after.lower()


def test_in_memory_zip_bundlers() -> None:
    z_name, z_bytes = create_humanized_docs_zip([("doc1.txt", "hello world"), ("doc2.md", b"# Title")])
    assert z_name.startswith("DraftProof_Humanized_Docs_") and z_name.endswith(".zip")
    with zipfile.ZipFile(io.BytesIO(z_bytes), "r") as zf:
        assert set(zf.namelist()) == {"doc1.txt", "doc2.md"}

    c_name, c_bytes = create_citations_zip("Cited text [1].", "@article{k,}", "\\bibitem{k}", "[{}]")
    assert c_name.startswith("DraftProof_Citations_")
    with zipfile.ZipFile(io.BytesIO(c_bytes), "r") as zf:
        assert "citations.bib" in zf.namelist()
        assert "citations.json" in zf.namelist()

    r_name, r_bytes = create_checker_reports_zip([("report.json", "{}")])
    assert r_name.startswith("DraftProof_Checker_Reports_")
